import io
import os
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_LABELS = {
    "TR": {
        "athletes": "Sporcuların Adı/Soyadı:",
        "team_no": "Takım NO:",
        "d_puani": "D- Puanı",
        "element_no": "Element NO",
        "coach_phone": "Antrenörün Telefon NO:",
        "club_name": "Kulübün Adı",
        "age_groups": "Yaş Grupları",
        "category": "Yarışma Kategorisi",
        "d_skoru": "D- Skoru",
        "cats": ["Çift Erkek", "Çift Kadın", "Çift Mix", "3'lü Kadın", "4'lü Erkek"]
    },
    "INT": {
        "athletes": "Gymnasts' Names:",
        "team_no": "Team NO:",
        "d_puani": "D- Value",
        "element_no": "Element NO",
        "coach_phone": "Coach Phone NO:",
        "club_name": "Club / Federation",
        "age_groups": "Age Groups",
        "category": "Competition Category",
        "d_skoru": "D- Score",
        "cats": ["Men's Pair", "Women's Pair", "Mixed Pair", "Women's Group", "Men's Group"]
    }
}

CAT_INDEX_MAP = {
    "Çift Erkek": 0,
    "Çift Kadın": 1,
    "Çift Mix": 2,
    "3'lü Kadın": 3,
    "4'lü Erkek": 4
}


def set_cell_margins(cell, top=20, bottom=20, left=30, right=30):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_text(cell, text, size=8, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    return p


def combine_images_vertically(img_paths, target_width=400, target_height=480):
    valid_imgs = [Image.open(p).convert("RGB") for p in img_paths if p and os.path.exists(p)]
    if not valid_imgs:
        return None

    canvas = Image.new("RGB", (target_width, target_height), (255, 255, 255))
    n = len(valid_imgs)
    slot_h = (target_height - 10 * (n + 1)) // n

    for idx, im in enumerate(valid_imgs):
        im_ratio = im.width / im.height
        max_w = target_width - 20
        max_h = slot_h

        if max_w / max_h > im_ratio:
            new_h = max_h
            new_w = int(new_h * im_ratio)
        else:
            new_w = max_w
            new_h = int(new_w / im_ratio)

        resized = im.resize((max(1, new_w), max(1, new_h)), Image.Resampling.LANCZOS)
        x_off = (target_width - new_w) // 2
        y_off = 10 + idx * (slot_h + 10) + (slot_h - new_h) // 2
        canvas.paste(resized, (x_off, y_off))

    buf = io.BytesIO()
    canvas.save(buf, format="PNG")
    buf.seek(0)
    return buf


def build_elements_row_table(doc, start_box, end_box, boxes_data, labels):
    num_boxes = end_box - start_box + 1
    total_cols = 1 + num_boxes * 4

    table = doc.add_table(rows=4, cols=total_cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"

    label_w = Cm(1.7)
    sub_w = Cm(0.68)

    for row in table.rows:
        row.cells[0].width = label_w
        for c in range(1, total_cols):
            row.cells[c].width = sub_w

    table.rows[0].height = Cm(0.45)
    table.rows[1].height = Cm(3.35)
    table.rows[2].height = Cm(0.50)
    table.rows[3].height = Cm(0.50)

    top_left = table.cell(0, 0).merge(table.cell(1, 0))
    set_cell_text(top_left, "", size=7)
    set_cell_text(table.cell(2, 0), labels["d_puani"], size=7.5, bold=True)
    set_cell_text(table.cell(3, 0), labels["element_no"], size=7.5, bold=True)

    for b_idx in range(num_boxes):
        box_num = start_box + b_idx
        c0 = 1 + b_idx * 4
        b_info = boxes_data.get(box_num, {"mode": "empty"})

        # Шапка бокса: слева в 1-й ячейке номер бокса, а в ячейках 2-4 (справа сбоку) по центру секунды
        num_cell = table.cell(0, c0)
        hold_cell = table.cell(0, c0 + 1).merge(table.cell(0, c0 + 3))
        set_cell_margins(num_cell, 10, 10, 15, 15)
        set_cell_margins(hold_cell, 10, 10, 15, 15)

        set_cell_text(num_cell, str(box_num), size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        hold_str = b_info.get("hold", "")
        set_cell_text(hold_cell, hold_str, size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Центр бокса (Картинка или разделитель //)
        img_cell = table.cell(1, c0).merge(table.cell(1, c0 + 3))
        set_cell_margins(img_cell, 15, 15, 15, 15)
        img_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        mode = b_info.get("mode", "empty")
        if mode == "separator":
            sep_sum = b_info.get("sep_sum", "")
            txt = f"// {sep_sum}" if sep_sum != "" else "//"
            set_cell_text(img_cell, txt, size=14, bold=True)
        elif mode in ("team", "individual"):
            img_paths = b_info.get("images", [])
            buf = combine_images_vertically(img_paths)
            if buf:
                p = img_cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run()
                run.add_picture(buf, width=Cm(2.45), height=Cm(3.1))

        # Строка D-Puanı (4 клеточки; для парного/группового балл стоит во 2-й клеточке)
        d_vals = b_info.get("d_cells", ["", "", "", ""])
        for sub_i in range(4):
            d_cell = table.cell(2, c0 + sub_i)
            set_cell_margins(d_cell, 10, 10, 10, 10)
            val_str = d_vals[sub_i] if sub_i < len(d_vals) else ""
            set_cell_text(d_cell, val_str, size=8, bold=True)

        # Строка Element NO
        el_cell = table.cell(3, c0).merge(table.cell(3, c0 + 3))
        set_cell_margins(el_cell, 10, 10, 15, 15)
        el_no_str = b_info.get("element_no", "")
        set_cell_text(el_cell, el_no_str, size=8, bold=False)

    return table


def generate_tariff_docx(meta, boxes_data, total_d_puani, total_d_skoru):
    doc = Document()
    comp_type = meta.get("comp_type", "TR")
    labels = DOC_LABELS.get(comp_type, DOC_LABELS["TR"])

    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.7)
    section.left_margin = Cm(0.9)
    section.right_margin = Cm(0.9)

    header_tbl = doc.add_table(rows=1, cols=2)
    header_tbl.style = "Table Grid"
    header_tbl.rows[0].height = Cm(0.75)
    header_tbl.cell(0, 0).width = Cm(20.8)
    header_tbl.cell(0, 1).width = Cm(7.0)

    athletes_str = ", ".join([a for a in meta.get("athletes", []) if a.strip()])
    set_cell_text(
        header_tbl.cell(0, 0),
        f"{labels['athletes']}   {athletes_str}",
        size=9.5,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT
    )
    set_cell_text(
        header_tbl.cell(0, 1),
        f"{labels['team_no']}   {meta.get('team_no', '')}",
        size=9.5,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT
    )

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(2)
    p_space.paragraph_format.space_after = Pt(2)

    layout_tbl = doc.add_table(rows=1, cols=2)
    layout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    left_container = layout_tbl.cell(0, 0)
    right_container = layout_tbl.cell(0, 1)
    left_container.width = Cm(21.0)
    right_container.width = Cm(6.8)

    left_container.paragraphs[0].paragraph_format.space_after = Pt(0)
    build_elements_row_table(left_container, 1, 7, boxes_data, labels)

    mid_p = left_container.add_paragraph()
    mid_p.paragraph_format.space_before = Pt(3)
    mid_p.paragraph_format.space_after = Pt(3)

    build_elements_row_table(left_container, 8, 14, boxes_data, labels)

    bot_p = left_container.add_paragraph()
    bot_p.paragraph_format.space_before = Pt(4)
    bot_p.paragraph_format.space_after = Pt(0)

    footer_tbl = left_container.add_table(rows=1, cols=3)
    footer_tbl.style = "Table Grid"
    footer_tbl.rows[0].height = Cm(0.65)
    footer_tbl.cell(0, 0).width = Cm(10.5)
    footer_tbl.cell(0, 1).width = Cm(5.1)
    footer_tbl.cell(0, 2).width = Cm(5.1)

    set_cell_text(
        footer_tbl.cell(0, 0),
        f"{labels['coach_phone']}  {meta.get('coach_phone', '')}",
        size=8.5,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT
    )
    set_cell_text(footer_tbl.cell(0, 1), "TF:", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(footer_tbl.cell(0, 2), "SR:", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    right_container.paragraphs[0].paragraph_format.space_after = Pt(0)

    club_tbl = right_container.add_table(rows=2, cols=1)
    club_tbl.style = "Table Grid"
    club_tbl.rows[0].height = Cm(0.5)
    club_tbl.rows[1].height = Cm(0.9)
    for r in club_tbl.rows:
        r.cells[0].width = Cm(6.5)
    set_cell_text(club_tbl.cell(0, 0), labels["club_name"], size=8.5, bold=True)
    set_cell_text(club_tbl.cell(1, 0), meta.get("club_name", ""), size=10, bold=True)

    p1 = right_container.add_paragraph()
    p1.paragraph_format.space_before = Pt(3)
    p1.paragraph_format.space_after = Pt(3)

    age_tbl = right_container.add_table(rows=3, cols=3)
    age_tbl.style = "Table Grid"
    age_tbl.rows[0].height = Cm(0.45)
    age_tbl.rows[1].height = Cm(0.45)
    age_tbl.rows[2].height = Cm(0.85)
    hdr_age = age_tbl.cell(0, 0).merge(age_tbl.cell(0, 2))
    set_cell_text(hdr_age, labels["age_groups"], size=8.5, bold=True)

    ages = ["6-12", "8-14", "9-15"]
    selected_age = meta.get("age_group", "6-12")
    for idx, ag in enumerate(ages):
        age_tbl.cell(1, idx).width = Cm(2.16)
        age_tbl.cell(2, idx).width = Cm(2.16)
        set_cell_text(age_tbl.cell(1, idx), ag, size=8, bold=True)
        mark = "V" if selected_age.startswith(ag) else ""
        set_cell_text(age_tbl.cell(2, idx), mark, size=14, bold=True)

    p2 = right_container.add_paragraph()
    p2.paragraph_format.space_before = Pt(3)
    p2.paragraph_format.space_after = Pt(3)

    cat_tbl = right_container.add_table(rows=3, cols=5)
    cat_tbl.style = "Table Grid"
    cat_tbl.rows[0].height = Cm(0.45)
    cat_tbl.rows[1].height = Cm(0.65)
    cat_tbl.rows[2].height = Cm(0.85)
    hdr_cat = cat_tbl.cell(0, 0).merge(cat_tbl.cell(0, 4))
    set_cell_text(hdr_cat, labels["category"], size=8.5, bold=True)

    selected_cat_idx = CAT_INDEX_MAP.get(meta.get("category", "Çift Erkek"), 0)
    for idx, ct_label in enumerate(labels["cats"]):
        cat_tbl.cell(1, idx).width = Cm(1.3)
        cat_tbl.cell(2, idx).width = Cm(1.3)
        set_cell_text(cat_tbl.cell(1, idx), ct_label, size=6.5, bold=True)
        mark = "V" if idx == selected_cat_idx else ""
        set_cell_text(cat_tbl.cell(2, idx), mark, size=13, bold=True)

    p3 = right_container.add_paragraph()
    p3.paragraph_format.space_before = Pt(4)
    p3.paragraph_format.space_after = Pt(4)

    score_tbl = right_container.add_table(rows=2, cols=2)
    score_tbl.style = "Table Grid"
    score_tbl.rows[0].height = Cm(0.5)
    score_tbl.rows[1].height = Cm(1.3)
    for c in range(2):
        score_tbl.cell(0, c).width = Cm(3.25)
        score_tbl.cell(1, c).width = Cm(3.25)

    set_cell_text(score_tbl.cell(0, 0), labels["d_skoru"], size=8.5, bold=True)
    set_cell_text(score_tbl.cell(0, 1), labels["d_puani"], size=8.5, bold=True)

    score_fmt = f"{total_d_skoru:.2f}".replace(".", ",") if comp_type == "TR" else f"{total_d_skoru:.2f}"
    set_cell_text(score_tbl.cell(1, 0), score_fmt, size=14, bold=True)
    set_cell_text(score_tbl.cell(1, 1), str(total_d_puani), size=14, bold=True)

    out_stream = io.BytesIO()
    doc.save(out_stream)
    out_stream.seek(0)
    return out_stream